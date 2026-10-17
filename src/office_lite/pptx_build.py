"""Build pptx bytes from blocks or csv rows (DRY, template-free)."""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Inches

from .guards import MAX_SLIDES, MAX_TABLE_COLS, check_size
from .md_parse import DocBlocks, TableBlock


def _new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs


def build_pptx(blocks: DocBlocks, title: str = "Deck") -> bytes:
    prs = _new_deck()
    layout_title = prs.slide_layouts[0]
    layout_content = prs.slide_layouts[1]
    first = prs.slides.add_slide(layout_title)
    first.shapes.title.text = title
    chunks: list[tuple[str, list[str]]] = []
    current_title = title
    current_points: list[str] = []
    for level, text in blocks.headings:
        if level <= 2:
            if current_points or current_title != title:
                chunks.append((current_title, current_points))
            current_title = text
            current_points = []
        else:
            current_points.append(text)
    for para in blocks.paragraphs:
        current_points.append(para)
    for _, item in blocks.bullets:
        current_points.append(item)
    chunks.append((current_title, current_points))
    planned = len(chunks) + len(blocks.tables) + len(blocks.codes) + 1
    if planned > MAX_SLIDES:
        raise ValueError("deck exceeds slide cap")
    for slide_title, points in chunks:
        slide = prs.slides.add_slide(layout_content)
        slide.shapes.title.text = slide_title
        body = slide.placeholders[1].text_frame
        body.clear()
        for point in points[:20]:
            para = body.add_paragraph()
            para.text = point
            para.level = 0
    for table in blocks.tables:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        _add_table(slide, table)
    for code in blocks.codes:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(6))
        box.text_frame.text = code[:2000]
    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    check_size(data)
    return data


def build_pptx_from_rows(header: list[str], rows: list[list[str]], title: str = "Table") -> bytes:
    prs = _new_deck()
    per_slide = 15
    pages = max((len(rows) + per_slide - 1) // per_slide, 1)
    if pages > MAX_SLIDES:
        raise ValueError("deck exceeds slide cap")
    for start in range(0, max(len(rows), 1), per_slide):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        if slide.shapes.title:
            slide.shapes.title.text = title
        window = rows[start : start + per_slide]
        _add_table(slide, TableBlock(header=header, rows=window))
        if not window:
            break
    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    check_size(data)
    return data


def _add_table(slide, table: TableBlock) -> None:
    header = table.header or (
        [f"col{i + 1}" for i in range(len(table.rows[0]))] if table.rows else []
    )
    if not header:
        return
    header = header[:MAX_TABLE_COLS]
    data_rows = [r[:MAX_TABLE_COLS] for r in table.rows]
    shape = slide.shapes.add_table(
        len(data_rows) + 1, len(header), Inches(0.5), Inches(1.0), Inches(12), Inches(5)
    )
    tbl = shape.table
    for i, name in enumerate(header):
        tbl.cell(0, i).text = name
    for r, row in enumerate(data_rows, start=1):
        padded = (row + [""] * len(header))[: len(header)]
        for i, value in enumerate(padded):
            tbl.cell(r, i).text = value
