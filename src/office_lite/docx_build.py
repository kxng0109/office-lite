"""Build docx bytes from blocks or csv rows (DRY, no template required)."""

from __future__ import annotations

import io

from docx import Document

from .guards import check_size
from .md_parse import DocBlocks, TableBlock


def build_docx(blocks: DocBlocks, title: str = "Document") -> bytes:
    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for level, text in blocks.headings:
        doc.add_heading(text, min(max(level, 1), 9))
    for para in blocks.paragraphs:
        doc.add_paragraph(para)
    for _, item in blocks.bullets:
        doc.add_paragraph(item, style="List Bullet")
    for code in blocks.codes:
        doc.add_paragraph(code, style="Normal")
    for table in blocks.tables:
        _add_table(doc, table)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    check_size(data)
    return data


def build_docx_from_rows(header: list[str], rows: list[list[str]], title: str = "Table") -> bytes:
    doc = Document()
    if title:
        doc.add_heading(title, 1)
    _add_table(doc, TableBlock(header=header, rows=rows))
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    check_size(data)
    return data


def _add_table(doc: Document, table: TableBlock) -> None:
    if not table.header and not table.rows:
        return
    header = (
        table.header or [f"col{i + 1}" for i in range(len(table.rows[0]))] if table.rows else []
    )
    if not header:
        return
    grid = doc.add_table(rows=1, cols=len(header))
    grid.style = "Light Grid Accent 1"
    for i, name in enumerate(header):
        grid.cell(0, i).text = name
    for row in table.rows:
        padded = (row + [""] * len(header))[: len(header)]
        cells = grid.add_row().cells
        for i, value in enumerate(padded):
            cells[i].text = value
