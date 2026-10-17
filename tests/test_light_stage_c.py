import pytest

from office_lite import guards
from office_lite.csv_bridge import rows_to_markdown
from office_lite.md_parse import get_parser
from office_lite.server import _walk_shapes, convert_markdown


def test_title_and_outline_caps():
    with pytest.raises(ValueError, match="title"):
        convert_markdown("# Hi", "docx", "")
    with pytest.raises(ValueError, match="title"):
        convert_markdown("# Hi", "docx", 123)  # type: ignore[arg-type]
    assert guards.check_outline([" a ", "b"]) == ["a", "b"]
    with pytest.raises(ValueError, match="non-empty"):
        guards.check_outline(["ok", "   "])


def test_char_budget_stops_early():
    header = ["h"]
    rows = [[f"row{i} with padding text"] for i in range(100)]
    full = rows_to_markdown(header, rows)
    capped = rows_to_markdown(header, rows, 200)
    assert len(capped) < len(full)
    assert capped.startswith("| h |")


def test_parser_singleton_shared():
    assert get_parser() is get_parser()


def test_walk_shapes_yields_plain_shapes():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    assert len(list(_walk_shapes(slide.shapes))) == len(list(slide.shapes))


def test_sanitize_coerces_non_string():
    assert guards.sanitize_csv_cell(123) == "123"
    assert guards.sanitize_csv_cell(None) == "None"


def test_path_jail_rejects_base_itself(tmp_path):
    with pytest.raises(ValueError, match="escapes"):
        guards.safe_output_path(tmp_path, ".")
