import base64
import io
import zipfile

import pytest

from office_lite import guards
from office_lite.csv_bridge import read_csv_text, rows_to_csv_text, rows_to_markdown
from office_lite.md_parse import parse_markdown, parse_text
from office_lite.server import (
    convert_markdown,
    convert_office_out,
    convert_table_in,
    create_from_outline,
    inspect_office,
    search,
)


def test_empty_and_whitespace_inputs():
    assert parse_markdown("").headings == []
    assert parse_text("   \n\n   ").paragraphs == []
    header, rows = read_csv_text("")
    assert header == [] and rows == []
    assert rows_to_markdown(["a"], []) == "| a |\n| --- |\n"
    assert "a" in rows_to_csv_text(["a"], [])


def test_special_characters_roundtrip():
    md = '# Caf\u00e9 \u4e2d\u6587 \U0001f600\n\nAT&T <tag> & "quotes"\n\n- a|b\n'
    blocks = parse_markdown(md)
    assert blocks.headings
    docx_b64 = convert_markdown(md, "docx", "S", "markdown")
    out = convert_office_out(docx_b64, "docx", "markdown")
    assert "Caf" in out
    pptx_b64 = convert_markdown(md, "pptx", "S", "markdown")
    assert "Slide" in convert_office_out(pptx_b64, "pptx", "text")


def test_oversized_rejected():
    big = "x" * (guards.MAX_INPUT_BYTES + 1)
    with pytest.raises(ValueError):
        convert_markdown(big, "docx")
    with pytest.raises(ValueError):
        convert_table_in(big, "markdown")
    with pytest.raises(ValueError):
        guards.check_size(b"x" * (guards.MAX_INPUT_BYTES + 1))


def test_malformed_b64_and_zip_rejected():
    with pytest.raises(ValueError, match="base64"):
        convert_office_out("!!!not-b64!!!", "docx", "markdown")
    with pytest.raises(ValueError, match="unreadable archive"):
        convert_office_out(base64.b64encode(b"not-a-zip").decode(), "docx", "markdown")
    with pytest.raises(ValueError, match="unreadable archive"):
        inspect_office(base64.b64encode(b"bad").decode(), "docx")


def test_csv_injection_sanitized_in_export():
    header = ["a"]
    rows = [["=1+1"], ["+2+2"], ["-3"], ["@sum"], ["\t=4"]]
    text = rows_to_csv_text(header, rows)
    assert "'=1+1" in text
    assert "'\t=4" in text
    header2, rows2 = read_csv_text("a;b\n1;2\n", ";")
    assert header2 == ["a", "b"] and rows2 == [["1", "2"]]


def test_csv_bom_and_crlf():
    text = "\ufeffa,b\r\n1,2\r\n"
    header, rows = read_csv_text(text)[:2]
    assert header == ["a", "b"]
    assert rows == [["1", "2"]]


def test_docx_merged_and_nested_tables():
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "merged"
    buf = io.BytesIO()
    doc.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    out = convert_office_out(b64, "docx", "markdown")
    assert "merged" in out
    assert "tables=1" in inspect_office(b64, "docx")


def test_pptx_tables_and_order():
    b64 = convert_table_in(
        "h1,h2\n" + "\n".join(f"r{i}c1,r{i}c2" for i in range(30)), "pptx", "Big"
    )
    info = inspect_office(b64, "pptx")
    assert info.startswith("slides=")
    assert int(info.split("slides=")[1].split()[0]) >= 2
    csv_out = convert_office_out(b64, "pptx", "csv")
    assert "h1" in csv_out


def test_path_jail_and_caps():
    import tempfile
    from pathlib import Path

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        guards.ensure_base_dir(base / "sub")
        assert (base / "sub").exists()
        with pytest.raises(ValueError):
            guards.safe_output_path(base, "../../evil.txt")
        with pytest.raises(ValueError):
            convert_table_in("x", "badtarget")


def test_malicious_markdown_html_not_executed():
    md = "<script>alert(1)</script>\n\n[click](javascript:alert(1))\n"
    b64 = convert_markdown(md, "docx", "XSS", "markdown")
    assert base64.b64decode(b64)[:2] == b"PK"
    b64p = convert_markdown(md, "pptx", "XSS", "markdown")
    assert base64.b64decode(b64p)[:2] == b"PK"


def test_zip_bomb_guard_entries():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(3):
            z.writestr(f"word/slide{i}.xml", "<a/>")
    assert zipfile.is_zipfile(io.BytesIO(buf.getvalue()))
    with pytest.raises(ValueError, match="unreadable archive"):
        convert_office_out(base64.b64encode(b"PK-not-office").decode(), "pptx", "text")


def test_outline_limits():
    many = [f"H{i}" for i in range(100)]
    b64 = create_from_outline(many, "pptx", "Big")
    assert base64.b64decode(b64)[:2] == b"PK"
    b64d = create_from_outline(many, "docx", "Big")
    assert base64.b64decode(b64d)[:2] == b"PK"


def test_search_bounds():
    assert search("docx", limit=1).count("\n") == 0
    assert search("zzz-no-match-xyz") == "no matches"
    long_q = "x" * 500
    assert isinstance(search(long_q), str)
