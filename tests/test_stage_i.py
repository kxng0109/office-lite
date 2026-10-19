import base64
import io

import pytest

from office_lite import guards
from office_lite.md_parse import DocBlocks, TableBlock
from office_lite.server import (
    _b64,
    _decode_office_b64,
    convert_markdown,
    convert_office_out,
    convert_table_in,
    inspect_office,
)


def test_strip_set_covers_known_smugglers():
    must_strip = [
        0x09,
        0x0B,
        0x0C,
        0x20,
        0x85,
        0xA0,
        0x1680,
        0x2000,
        0x2001,
        0x2002,
        0x2003,
        0x2004,
        0x2005,
        0x2006,
        0x2007,
        0x2008,
        0x2009,
        0x200A,
        0x2028,
        0x2029,
        0x202F,
        0x205F,
        0x3000,
        0x200B,
        0x200C,
        0x200D,
        0xFEFF,
        0x180E,
        0x2060,
        0x00AD,
        0x3164,
        0xFFA0,
    ]
    stripped = set(guards._STRIP_LEAD)
    for code in must_strip:
        assert chr(code) in stripped, hex(code)
    assert guards.sanitize_csv_cell("\u2028=cmd") == "'\u2028=cmd"
    assert guards.sanitize_csv_cell("\u00a0+1") == "'\u00a0+1"


def test_strict_input_types():
    with pytest.raises(ValueError, match="must be a string"):
        convert_markdown(123, "docx")
    with pytest.raises(ValueError, match="must be a string"):
        convert_table_in(None, "markdown")
    with pytest.raises(ValueError, match="must be a string"):
        _decode_office_b64(123)
    with pytest.raises(ValueError, match="must be a string"):
        inspect_office(None, "docx")


def test_nested_table_depth_cap():
    from types import SimpleNamespace

    from office_lite.server import _walk_tables

    def nest(depth):
        if depth == 0:
            return SimpleNamespace(rows=[])
        cell = SimpleNamespace(tables=[nest(depth - 1)])
        return SimpleNamespace(rows=[SimpleNamespace(cells=[cell])])

    with pytest.raises(ValueError, match="depth cap"):
        list(_walk_tables([nest(guards.MAX_GROUP_DEPTH + 2)]))


def test_deeply_nested_docx_fails_closed():
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    current = table.cell(0, 0)
    for _ in range(guards.MAX_TABLES):
        inner = current.add_table(rows=1, cols=1)
        current = inner.cell(0, 0)
    buf = io.BytesIO()
    doc.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    with pytest.raises(ValueError):
        convert_office_out(b64, "docx", "markdown")


def test_table_count_cap_enforced_on_build():
    blocks = DocBlocks(
        headings=[],
        tables=[TableBlock(header=["h"], rows=[["v"]]) for _ in range(guards.MAX_TABLES + 1)],
    )
    with pytest.raises(ValueError, match="table cap"):
        from office_lite.docx_build import build_docx

        build_docx(blocks)


def test_total_cells_cap_enforced():
    from office_lite.docx_build import build_docx

    blocks = DocBlocks(
        headings=[],
        tables=[TableBlock(header=["h"] * 100, rows=[["v"] * 100 for _ in range(700)])] * 3,
    )
    with pytest.raises(ValueError, match="total cell cap"):
        build_docx(blocks)


def test_column_cap_enforced_on_build():
    from office_lite.docx_build import build_docx_from_rows

    with pytest.raises(ValueError, match="column cap"):
        build_docx_from_rows([f"c{i}" for i in range(guards.MAX_TABLE_COLS + 1)], [["x"]])


def test_main_failure_exits_nonzero(monkeypatch):
    import office_lite.server as srv

    def boom():
        raise RuntimeError("loop dead")

    monkeypatch.setattr(srv.server, "run", boom)
    with pytest.raises(SystemExit) as exc:
        srv.main()
    assert exc.value.code == 1


def test_template_cache_reused():
    from office_lite.docx_build import _new_doc
    from office_lite.pptx_build import _new_deck

    assert _new_doc() is not _new_doc()
    assert _new_deck() is not _new_deck()
    import office_lite.docx_build as db
    import office_lite.pptx_build as pb

    assert db._cached_template is not None
    assert pb._cached_template is not None


def test_b64_output_bounded_by_construction():
    assert _b64(b"abc") == "YWJj"
