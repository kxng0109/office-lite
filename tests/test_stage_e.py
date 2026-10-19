import base64
import io
import threading

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from office_lite import guards
from office_lite.csv_bridge import read_csv_text
from office_lite.md_parse import parse_markdown
from office_lite.pptx_build import build_pptx
from office_lite.server import (
    _b64,
    _walk_shapes,
    convert_table_in,
    inspect_office,
)


def test_empty_csv_rejected():
    with pytest.raises(ValueError, match="empty"):
        convert_table_in("", "markdown")
    with pytest.raises(ValueError, match="empty"):
        convert_table_in("   \n", "csv")


def test_delimiter_allowlist():
    with pytest.raises(ValueError, match="delimiter"):
        read_csv_text("a,b", '"')
    with pytest.raises(ValueError, match="delimiter"):
        read_csv_text("a,b", "\\")
    header, rows = read_csv_text("a|b\n1|2", "|")
    assert header == ["a", "b"] and rows == [["1", "2"]]


def test_b64_output_cap():
    with pytest.raises(ValueError, match="size cap"):
        _b64(b"x" * (guards.MAX_INPUT_BYTES + 1))


def test_group_depth_cap():
    from types import SimpleNamespace

    def nest(depth):
        if depth == 0:
            return SimpleNamespace(shape_type=1, shapes=[])
        return SimpleNamespace(shape_type=MSO_SHAPE_TYPE.GROUP, shapes=[nest(depth - 1)])

    with pytest.raises(ValueError, match="depth cap"):
        list(_walk_shapes([nest(guards.MAX_GROUP_DEPTH + 2)]))


def test_large_table_paginated_and_bounded():
    md = "| c1 | c2 |\n|---|---|\n" + "".join(f"| r{i} | v{i} |\n" for i in range(2000))
    blocks = parse_markdown(md)
    assert len(blocks.tables[0].rows) == 2000
    data = build_pptx(blocks, "Big")
    assert data[:2] == b"PK"
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) > 100


def test_threaded_parse_smoke():
    errors: list = []

    def work():
        try:
            parse_markdown("# T\n\n- a\n- b\n")
        except Exception as exc:  # noqa: BLE001
            errors.append(exc)

    threads = [threading.Thread(target=work) for _ in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors


def test_max_cells_guard():
    header = [f"c{i}" for i in range(100)]
    rows = [["x"] * 100 for _ in range(600)]
    with pytest.raises(ValueError, match="cell cap"):
        convert_table_in(",".join(header) + "\n" + "\n".join(",".join(r) for r in rows), "docx")


def test_inspect_rejects_bad_source_and_bytes():
    with pytest.raises(ValueError, match="source"):
        inspect_office(base64.b64encode(b"PK").decode(), "xlsx")
    with pytest.raises(ValueError, match="unreadable"):
        inspect_office(base64.b64encode(b"not-a-zip-at-all").decode(), "pptx")
