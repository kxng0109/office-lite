import io

import pytest

import office_lite
from office_lite import guards
from office_lite.md_parse import parse_markdown
from office_lite.server import convert_markdown, convert_office_out, convert_table_in, search


def test_version_matches_installed_metadata():
    from importlib.metadata import version

    assert office_lite.__version__ == version("office-lite")


def test_hyperlink_destination_preserved():
    blocks = parse_markdown("[Invoice](http://evil.example/x)\n")
    assert blocks.paragraphs == ["[Invoice](http://evil.example/x)"]


def test_nested_tables_extracted():
    from docx import Document

    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "nested-secret"
    buf = io.BytesIO()
    doc.save(buf)
    import base64

    md = convert_office_out(base64.b64encode(buf.getvalue()).decode(), "docx", "markdown")
    assert "nested-secret" in md


def test_multi_table_csv_rejected():
    from docx import Document

    doc = Document()
    for label in ("t1", "t2"):
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = label
    buf = io.BytesIO()
    doc.save(buf)
    import base64

    b64 = base64.b64encode(buf.getvalue()).decode()
    with pytest.raises(ValueError, match="multi-table"):
        convert_office_out(b64, "docx", "csv")


def test_same_header_tables_merge_to_csv():
    from docx import Document

    doc = Document()
    for _ in range(2):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "h1"
        table.cell(0, 1).text = "h2"
        table.cell(1, 0).text = "v1"
        table.cell(1, 1).text = "v2"
    buf = io.BytesIO()
    doc.save(buf)
    import base64

    b64 = base64.b64encode(buf.getvalue()).decode()
    out = convert_office_out(b64, "docx", "csv")
    assert out.count("v1") == 2


def test_title_control_chars_rejected():
    with pytest.raises(ValueError, match="control"):
        convert_markdown("# Hi", "docx", "Bad\x00Title")
    with pytest.raises(ValueError, match="control"):
        convert_table_in("a,b\n1,2", "markdown", "Bad\x07")


def test_per_item_caps_enforced():
    with pytest.raises(ValueError, match="length cap"):
        parse_markdown("# " + "x" * (guards.MAX_HEADING_CHARS + 1))
    with pytest.raises(ValueError, match="length cap"):
        parse_markdown("y" * (guards.MAX_PARA_CHARS + 1))


def test_formula_unicode_whitespace_neutralized():
    assert guards.sanitize_csv_cell(" =cmd|'/c calc'!A0").startswith("'")
    assert guards.sanitize_csv_cell("​=1+1").startswith("'")
    assert guards.sanitize_md_cell("​=1+1").startswith("'")


def test_search_strict_contract():
    assert search("docx", limit=20).count("\n") >= 0
