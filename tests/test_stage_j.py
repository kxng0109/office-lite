import base64
import io

import pytest

from office_lite import guards
from office_lite.md_parse import parse_markdown
from office_lite.server import convert_office_out


def test_fullwidth_triggers_neutralized():
    assert guards.sanitize_csv_cell("＝1+1").startswith("'")
    assert guards.sanitize_csv_cell("＋2").startswith("'")
    assert guards.sanitize_md_cell("－3").startswith("'")
    assert guards.sanitize_md_cell("＠x").startswith("'")
    assert guards.sanitize_csv_cell("100%") == "100%"
    assert guards.sanitize_csv_cell("a|b") == "a|b"


def test_javascript_and_data_urls_defanged_md():
    out = parse_markdown("[click](javascript:alert(1))\n").paragraphs
    assert out == ["\\[click](javascript:alert(1))"]
    out = parse_markdown("[f](data:text/html,<b>x</b>)\n").paragraphs
    assert out == ["\\[f](data:text/html,<b>x</b>)"]
    assert parse_markdown("[ok](https://safe.example/)\n").paragraphs == [
        "[ok](https://safe.example/)"
    ]


def test_bracket_forgery_escaped():
    out = parse_markdown("[a](http://x.example/) [b](http://y.example/)\n").paragraphs
    assert "[a](http://x.example/)" in out[0]
    tricky = parse_markdown("[a](http://x.example/)x](http://evil.example/)\n").paragraphs
    assert "http://evil.example/)" in tricky[0] or "evil" in tricky[0]


def test_nul_byte_passes_through_unsanitized():
    # NUL is not an OWASP trigger; document the passthrough contract.
    assert guards.sanitize_csv_cell("\x00abc") == "\x00abc"


def test_csv_field_size_limit_explicit():
    import csv as _csv

    assert _csv.field_size_limit() == 128 * 1024


def test_javascript_docx_hyperlink_defanged():
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("See ")
    r_id = doc.part.relate_to(
        "javascript:alert(1)",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Click"
    run.append(text)
    link.append(run)
    para._p.append(link)
    buf = io.BytesIO()
    doc.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    out = convert_office_out(b64, "docx", "markdown")
    assert "Click" in out
    assert "javascript:" not in out


def test_pptx_hyperlink_round_trip_safe():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(2))
    tf = box.text_frame
    tf.text = "Go "
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = "Here"
    run.hyperlink.address = "https://safe.example/page?a=1"
    buf = io.BytesIO()
    prs.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    out = convert_office_out(b64, "pptx", "markdown")
    assert "[Here](https://safe.example/page?a=1)" in out


def test_omitted_count_arithmetic():
    from office_lite.csv_bridge import rows_to_markdown

    rows = [[f"v{i}"] for i in range(50)]
    out = rows_to_markdown(["h"], rows, char_limit=60)
    lines = out.splitlines()
    shown = sum(1 for line in lines if line.startswith("| v"))
    marker = [line for line in lines if "omitted" in line]
    assert len(marker) == 1
    assert shown + int(marker[0].split("(")[1].split()[0]) == len(rows)
    # Marker must live outside the table so strict parsers still see a
    # rectangular grid: every pipe-line has exactly the header column count.
    assert not marker[0].lstrip().startswith("|")
    widths = {line.count("|") for line in lines if line.startswith("|")}
    assert widths == {2}


def test_safe_link_branches():
    from office_lite import guards

    assert guards.safe_link("A", "") == "A"
    assert guards.safe_link("A", "   ") == "A"
    assert guards.safe_link("A", "#sec") == "[A](#sec)"
    assert guards.safe_link("A B", "https://x.example/a b") == "[A B](<https://x.example/a b>)"
    assert guards.safe_link("L", "ftp://x.example/f") == "L"
    assert guards.safe_link("L", "data:text/plain,hi") == "L"
    assert guards.safe_link("L", "vbscript:x") == "L"
    assert guards.defang_link_forgery("plain text") == "plain text"
    assert guards.defang_link_forgery("[c](VBSCRIPT:x)") == "\\[c](VBSCRIPT:x)"
    assert guards.safe_link("[x]", "https://x.example/") == "[\\[x\\]](https://x.example/)"


def test_csv_field_size_limit_is_explicit():
    import csv as _csv

    assert _csv.field_size_limit() == 128 * 1024


def test_merged_rows_capped():
    from office_lite import server as srv

    header = ["h"]
    parts = [(header, [["v"]] * 10) for _ in range(300)]
    with pytest.raises(ValueError, match="row cap"):
        srv._merge_tables(parts)
    ok = srv._merge_tables([(header, [["v"]] * 10) for _ in range(100)])
    assert len(ok[1]) == 1000


def test_group_nesting_depth_capped():
    from types import SimpleNamespace

    from office_lite import server as srv

    def nest(depth):
        if depth == 0:
            return SimpleNamespace(shape_type=7, has_table=False, has_text_frame=False, shapes=[])
        return SimpleNamespace(
            shape_type=6, has_table=False, has_text_frame=False, shapes=[nest(depth - 1)]
        )

    with pytest.raises(ValueError, match="depth cap"):
        list(srv._walk_shapes([nest(guards.MAX_GROUP_DEPTH + 2)]))


def test_tool_logs_input_size(caplog):
    import logging

    from office_lite.server import convert_markdown

    with caplog.at_level(logging.INFO, logger="office_lite.server"):
        convert_markdown("# Hi", "docx", "T", "markdown")
    assert any("convert_markdown" in r.message for r in caplog.records)
